"""
generate_procedure.py
Kurumsal kimlige uygun, A3 yatay standart montaj proseduru (.docx) uretici.

Kullanim:
    python generate_procedure.py --urun urunler/POMPA-BLOK-260 \
        --config config/kurumsal_kimlik.json --out cikti/
"""

import json
import csv
import argparse
from pathlib import Path

import yaml
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.shared import RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SAYFA_GENISLIK_CM = 42.0  # A3 yatay
SAYFA_YUKSEKLIK_CM = 29.7
YAN_MARJ_CM = 2.0
ICERIK_GENISLIK_CM = SAYFA_GENISLIK_CM - 2 * YAN_MARJ_CM  # 38.0 cm
HEADER_MESAFE_CM = 1.0
FOOTER_MESAFE_CM = 1.0
# Header/footer artik metin+logo olarak kuruluyor (banner gorseli basilmiyor);
# bu degerler o metin bloklarinin tahmini yuksekligine gore ayarlanmis tampon paylardir.
UST_MARJ_CM = 3.2
ALT_MARJ_CM = 3.6
LOGO_GENISLIK_CM = 4.0


def load_kurumsal_kimlik(config_path: Path) -> dict:
    with open(config_path, encoding="utf-8") as f:
        return json.load(f)


def load_urun_bilgisi(urun_dir: Path) -> dict:
    with open(urun_dir / "urun_bilgisi.json", encoding="utf-8") as f:
        return json.load(f)


def load_malzeme_listesi(urun_dir: Path) -> list:
    with open(urun_dir / "malzeme_listesi.csv", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load_islem_adimlari(urun_dir: Path) -> list:
    with open(urun_dir / "islem_adimlari.yaml", encoding="utf-8") as f:
        return yaml.safe_load(f)["adimlar"]


def hex_to_rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_paragraph_rule(paragraph, kenar: str, color_hex: str, boyut: int = 6):
    """Paragrafin ustune/altina ince bir cizgi (border) ekler - kurumsal header/footer ayirici."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    kenar_el = OxmlElement(f"w:{kenar}")
    kenar_el.set(qn("w:val"), "single")
    kenar_el.set(qn("w:sz"), str(boyut))
    kenar_el.set(qn("w:space"), "4")
    kenar_el.set(qn("w:color"), color_hex.lstrip("#"))
    pBdr.append(kenar_el)
    pPr.append(pBdr)


def add_header(doc: Document, kimlik: dict):
    section = doc.sections[0]
    header = section.header
    icerik_genislik = section.page_width - section.left_margin - section.right_margin
    ana_renk = kimlik.get("ana_renk_hex", "#CC0605")
    ikincil_renk = kimlik.get("ikincil_renk_hex", "#9B9B94")

    # 1. satir: logo (sol) + kurum adi (sag) - tek satirda tab ile hizalanir
    satir1 = header.paragraphs[0]
    satir1.paragraph_format.space_before = Pt(0)
    satir1.paragraph_format.space_after = Pt(0)
    satir1.paragraph_format.tab_stops.add_tab_stop(icerik_genislik, WD_TAB_ALIGNMENT.RIGHT)

    if kimlik.get("logo_yolu") and Path(kimlik["logo_yolu"]).exists():
        logo_run = satir1.add_run()
        logo_run.add_picture(kimlik["logo_yolu"], width=Cm(LOGO_GENISLIK_CM))

    satir1.add_run("\t")
    firma_run = satir1.add_run(kimlik.get("firma_adi", ""))
    firma_run.bold = True
    firma_run.font.size = Pt(16)
    firma_run.font.color.rgb = hex_to_rgb(ana_renk)

    # 2. satir: slogan (sag), altinda ayirici cizgi
    if kimlik.get("slogan"):
        satir2 = header.add_paragraph()
        satir2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        satir2.paragraph_format.space_before = Pt(2)
        satir2.paragraph_format.space_after = Pt(0)
        slogan_run = satir2.add_run(kimlik["slogan"])
        slogan_run.font.size = Pt(10)
        slogan_run.font.color.rgb = hex_to_rgb(ikincil_renk)
        set_paragraph_rule(satir2, "bottom", ana_renk)
    else:
        set_paragraph_rule(satir1, "bottom", ana_renk)


def _footer_satiri(footer, sol_metin, sag_metin, icerik_genislik, kalin=False, sol_pt=9, sag_pt=8, sol_renk=None, sag_renk=None, yeni=True):
    p = footer.add_paragraph() if yeni else footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(icerik_genislik, WD_TAB_ALIGNMENT.RIGHT)
    if sol_metin:
        r = p.add_run(sol_metin)
        r.bold = kalin
        r.font.size = Pt(sol_pt)
        if sol_renk:
            r.font.color.rgb = hex_to_rgb(sol_renk)
    if sag_metin:
        p.add_run("\t")
        r2 = p.add_run(sag_metin)
        r2.font.size = Pt(sag_pt)
        if sag_renk:
            r2.font.color.rgb = hex_to_rgb(sag_renk)
    return p


def add_footer(doc: Document, kimlik: dict, urun: dict):
    section = doc.sections[0]
    footer = section.footer
    icerik_genislik = section.page_width - section.left_margin - section.right_margin
    ana_renk = kimlik.get("ana_renk_hex", "#CC0605")
    ikincil_renk = kimlik.get("ikincil_renk_hex", "#9B9B94")

    iletisim_satiri = " • ".join(
        x for x in (kimlik.get("web"), kimlik.get("eposta")) if x
    )
    tel_satiri = " • ".join(
        x for x in (
            f"Tel: {kimlik['telefon']}" if kimlik.get("telefon") else None,
            f"Faks: {kimlik['faks']}" if kimlik.get("faks") else None,
        ) if x
    )

    satir1 = _footer_satiri(
        footer, kimlik.get("firma_adi", ""), kimlik.get("vergi_dairesi", ""),
        icerik_genislik, kalin=True, sol_pt=9, sag_pt=7.5,
        sol_renk=ana_renk, sag_renk=ikincil_renk, yeni=False,
    )
    set_paragraph_rule(satir1, "top", ana_renk)

    vergi_no = f"Vergi Numarası: {kimlik['vergi_no']}" if kimlik.get("vergi_no") else ""
    _footer_satiri(footer, kimlik.get("adres", ""), vergi_no, icerik_genislik, sag_renk=ikincil_renk)

    mersis_no = f"Mersis No: {kimlik['mersis_no']}" if kimlik.get("mersis_no") else ""
    _footer_satiri(footer, iletisim_satiri, mersis_no, icerik_genislik, sag_renk=ikincil_renk)

    ticaret_sicil = f"Ticaret Sicil No: {kimlik['ticaret_sicil_no']}" if kimlik.get("ticaret_sicil_no") else ""
    _footer_satiri(footer, tel_satiri, ticaret_sicil, icerik_genislik, sag_renk=ikincil_renk)

    p_kontrol = footer.add_paragraph()
    p_kontrol.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kontrol.paragraph_format.space_before = Pt(4)
    p_kontrol.paragraph_format.space_after = Pt(0)
    kontrol_run = p_kontrol.add_run(
        f"Dok No: {urun.get('dok_no', '')} | Rev: {urun.get('revizyon', '')} | "
        f"Tarih: {urun.get('tarih', '')}"
    )
    kontrol_run.font.size = Pt(8)
    kontrol_run.font.color.rgb = hex_to_rgb(ikincil_renk)


def set_document_font(doc: Document, kimlik: dict):
    font_adi = kimlik.get("font_adi", "Calibri")
    for style_adi in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_adi]
        style.font.name = font_adi
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), font_adi)
        rfonts.set(qn("w:cs"), font_adi)


def set_page_layout(doc: Document):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(SAYFA_GENISLIK_CM)
    section.page_height = Cm(SAYFA_YUKSEKLIK_CM)
    section.left_margin = Cm(YAN_MARJ_CM)
    section.right_margin = Cm(YAN_MARJ_CM)
    section.top_margin = Cm(UST_MARJ_CM)
    section.bottom_margin = Cm(ALT_MARJ_CM)
    section.header_distance = Cm(HEADER_MESAFE_CM)
    section.footer_distance = Cm(FOOTER_MESAFE_CM)


def add_baslik(doc: Document, kimlik: dict, urun: dict):
    baslik = doc.add_heading(f"{urun['urun_adi']} — Montaj Prosedürü", level=1)
    if kimlik.get("ana_renk_hex"):
        for run in baslik.runs:
            run.font.color.rgb = hex_to_rgb(kimlik["ana_renk_hex"])


def add_patlamis_gorunum(doc: Document, urun_dir: Path, urun: dict):
    doc.add_heading("1. Patlamış Görünüm ve Malzeme Listesi", level=2)
    gorsel_yolu = urun_dir / urun.get("patlamis_gorunum_dosya", "patlamis_gorunum.png")
    if gorsel_yolu.exists():
        doc.add_picture(str(gorsel_yolu), width=Cm(24))
    else:
        doc.add_paragraph(f"[Görsel bulunamadı: {gorsel_yolu.name}]")


def add_malzeme_tablosu(doc: Document, malzemeler: list):
    if not malzemeler:
        return
    kolonlar = list(malzemeler[0].keys())
    tablo = doc.add_table(rows=1, cols=len(kolonlar))
    tablo.style = "Light Grid Accent 1"
    tablo.alignment = WD_TABLE_ALIGNMENT.CENTER
    tablo.autofit = False
    kolon_genislik = Cm(ICERIK_GENISLIK_CM / len(kolonlar))
    hdr_cells = tablo.rows[0].cells
    for i, kolon in enumerate(kolonlar):
        hdr_cells[i].text = kolon
        hdr_cells[i].width = kolon_genislik
    for satir in malzemeler:
        row_cells = tablo.add_row().cells
        for i, kolon in enumerate(kolonlar):
            row_cells[i].text = str(satir.get(kolon, ""))
            row_cells[i].width = kolon_genislik
    for kolon in tablo.columns:
        kolon.width = kolon_genislik


def add_islem_adimlari(doc: Document, adimlar: list):
    doc.add_heading("2. Montaj Adımları", level=2)
    for adim in adimlar:
        doc.add_heading(f"Adım {adim['no']}: {adim['baslik']}", level=3)
        doc.add_paragraph(adim["aciklama"])
        if adim.get("tork"):
            doc.add_paragraph(f"Sıkma torku: {adim['tork']}", style="Intense Quote")
        if adim.get("gorsel"):
            gorsel_yolu = Path(adim["gorsel"])
            if gorsel_yolu.exists():
                doc.add_picture(str(gorsel_yolu), width=Cm(14))


def generate(urun_dir: Path, config_path: Path, out_dir: Path):
    kimlik = load_kurumsal_kimlik(config_path)
    urun = load_urun_bilgisi(urun_dir)
    malzemeler = load_malzeme_listesi(urun_dir)
    adimlar = load_islem_adimlari(urun_dir)

    doc = Document()
    set_page_layout(doc)
    set_document_font(doc, kimlik)
    add_header(doc, kimlik)
    add_footer(doc, kimlik, urun)
    add_baslik(doc, kimlik, urun)
    add_patlamis_gorunum(doc, urun_dir, urun)
    add_malzeme_tablosu(doc, malzemeler)
    add_islem_adimlari(doc, adimlar)

    out_dir.mkdir(parents=True, exist_ok=True)
    cikti_adi = f"{urun['urun_kodu']}_Rev{urun.get('revizyon', 'A')}_Montaj_Proseduru.docx"
    doc.save(out_dir / cikti_adi)
    print(f"Oluşturuldu: {out_dir / cikti_adi}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--urun", required=True, type=Path)
    parser.add_argument("--config", default=Path("config/kurumsal_kimlik.json"), type=Path)
    parser.add_argument("--out", default=Path("cikti"), type=Path)
    args = parser.parse_args()
    generate(args.urun, args.config, args.out)
